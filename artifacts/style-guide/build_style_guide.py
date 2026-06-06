from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = Path(__file__).resolve().parent
OUTPUT = OUT_DIR / "Zap-It-Style-Guide.docx"
LOGO = OUT_DIR / "zap-it-logo.png"
SHAPES_IMAGE = OUT_DIR / "shape-language.png"

COLORS = {
    "ink": "17212B",
    "muted": "5F6D78",
    "panel": "FFFFFF",
    "line": "D8E1E7",
    "field": "EDF3F7",
    "accent": "0F7B8F",
    "accent_strong": "095B6D",
    "danger": "BC3152",
    "good": "16704A",
    "yellow": "FFC940",
    "yellow_light": "FFF16A",
    "orange": "EF8737",
    "blue": "2877D8",
    "green": "22A06B",
    "red": "D83B4A",
    "soft_teal": "DFFBFF",
    "timer": "F7FDFF",
    "dark_panel": "18242D",
    "dark_field": "22333E",
    "dark_line": "344955",
}


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_font(run, size=None, color=None, bold=None, italic=None, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D8E1E7", size=8):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_font(run, size=8.5, color=COLORS["muted"], bold=True)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_picture_with_alt(run, path, width, alt_text):
    shape = run.add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    return shape


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def add_text(doc, text, size=10.2, color=None, bold=False, italic=False, after=6, before=0,
             align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, size=size, color=color or COLORS["ink"], bold=bold, italic=italic)
    if keep:
        set_keep_with_next(p)
    return p


def add_rich_text(doc, parts, size=10.2, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    for text, bold, color in parts:
        run = p.add_run(text)
        set_font(run, size=size, color=color or COLORS["ink"], bold=bold)
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    set_font(r, size=8.5, color=COLORS["accent"], bold=True)
    r.font.letter_spacing = Pt(0.6)
    set_keep_with_next(p)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_keep_with_next(p)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=9.9, color=COLORS["ink"], bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, size=9.9, color=COLORS["ink"])
    else:
        r = p.add_run(text)
        set_font(r, size=9.9, color=COLORS["ink"])
    return p


def add_callout(doc, label, text, fill="EDF3F7", accent="0F7B8F"):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [260, 9100], 140)
    left, body = table.rows[0].cells
    set_cell_shading(left, accent)
    set_cell_shading(body, fill)
    for cell in (left, body):
        set_cell_margins(cell, top=120, bottom=120, start=140, end=140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = body.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label.upper()}  ")
    set_font(r, size=8.6, color=accent, bold=True)
    r = p.add_run(text)
    set_font(r, size=9.4, color=COLORS["ink"])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_standard_table(doc, headers, rows, widths, header_fill="0F7B8F", header_text="FFFFFF",
                       font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, 120)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, header_fill)
        set_cell_margins(cell, top=110, bottom=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, size=8.3, color=header_text, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else "F7FAFB")
            set_cell_border(cell, COLORS["line"], size=6)
            set_cell_margins(cell, top=105, bottom=105)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_font(r, size=font_size, color=COLORS["ink"], bold=(index == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_color_table(doc, entries, widths=(2250, 1300, 5810)):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, list(widths), 120)
    headers = ["Token", "Value", "Role"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, COLORS["ink"])
        set_cell_border(cell, COLORS["ink"])
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(header)
        set_font(r, size=8.3, color="FFFFFF", bold=True)
    set_repeat_table_header(table.rows[0])
    for name, value, role in entries:
        cells = table.add_row().cells
        set_cell_shading(cells[0], value)
        sample_text = "FFFFFF" if value in {
            COLORS["ink"], COLORS["accent"], COLORS["accent_strong"], COLORS["danger"],
            COLORS["good"], COLORS["dark_panel"], COLORS["dark_field"], COLORS["dark_line"],
            COLORS["blue"], COLORS["green"], COLORS["red"], COLORS["orange"],
        } else COLORS["ink"]
        r = cells[0].paragraphs[0].add_run(name)
        set_font(r, size=8.6, color=sample_text, bold=True)
        r = cells[1].paragraphs[0].add_run(f"#{value}")
        set_font(r, size=8.5, color=COLORS["ink"], bold=True)
        r = cells[2].paragraphs[0].add_run(role)
        set_font(r, size=8.5, color=COLORS["ink"])
        for cell in cells:
            set_cell_border(cell, COLORS["line"], size=6)
            set_cell_margins(cell, top=110, bottom=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def create_shape_image():
    width, height = 1300, 260
    image = Image.new("RGB", (width, height), "#F7FAFB")
    draw = ImageDraw.Draw(image)
    centers = [(130, 130), (390, 130), (650, 130), (910, 130), (1170, 130)]
    fills = ["#F0C83B", "#2877D8", "#22A06B", "#EF8737", "#D83B4A"]
    outline = "#66737C"
    draw.regular_polygon((centers[0][0], centers[0][1], 82), 10, rotation=18, fill=fills[0], outline=outline, width=6)
    draw.ellipse((centers[1][0] - 72, centers[1][1] - 72, centers[1][0] + 72, centers[1][1] + 72),
                 fill=fills[1], outline=outline, width=6)
    draw.rounded_rectangle((centers[2][0] - 72, centers[2][1] - 72, centers[2][0] + 72, centers[2][1] + 72),
                           radius=14, fill=fills[2], outline=outline, width=6)
    draw.polygon([(centers[3][0], centers[3][1] - 82), (centers[3][0] + 84, centers[3][1] + 72),
                  (centers[3][0] - 84, centers[3][1] + 72)], fill=fills[3], outline=outline)
    draw.line([(centers[3][0], centers[3][1] - 82), (centers[3][0] + 84, centers[3][1] + 72),
               (centers[3][0] - 84, centers[3][1] + 72), (centers[3][0], centers[3][1] - 82)],
              fill=outline, width=6, joint="curve")
    draw.polygon([(centers[4][0], centers[4][1] - 86), (centers[4][0] + 86, centers[4][1]),
                  (centers[4][0], centers[4][1] + 86), (centers[4][0] - 86, centers[4][1])],
                 fill=fills[4], outline=outline)
    draw.line([(centers[4][0], centers[4][1] - 86), (centers[4][0] + 86, centers[4][1]),
               (centers[4][0], centers[4][1] + 86), (centers[4][0] - 86, centers[4][1]),
               (centers[4][0], centers[4][1] - 86)], fill=outline, width=6, joint="curve")
    image.save(SHAPES_IMAGE)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(COLORS["ink"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    heading_specs = {
        "Heading 1": (18, COLORS["ink"], 16, 7),
        "Heading 2": (13, COLORS["accent"], 11, 5),
        "Heading 3": (10.5, COLORS["accent_strong"], 8, 3),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = doc.styles["List Bullet"]
    list_style.font.name = "Arial"
    list_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    list_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    list_style.font.size = Pt(9.9)
    list_style.paragraph_format.left_indent = Inches(0.38)
    list_style.paragraph_format.first_line_indent = Inches(-0.19)
    list_style.paragraph_format.space_after = Pt(4)
    list_style.paragraph_format.line_spacing = 1.18

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    table = header.add_table(rows=1, cols=2, width=Inches(6.78))
    set_table_geometry(table, [6800, 2560], 0)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=40, start=0, end=0)
    r = left.paragraphs[0].add_run("ZAP IT  /  STYLE GUIDE")
    set_font(r, size=8, color=COLORS["accent"], bold=True)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = right.paragraphs[0].add_run("Product design reference")
    set_font(r, size=8, color=COLORS["muted"])
    for cell in (left, right):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:color"), COLORS["line"])
        borders.append(bottom)
        tc_pr.append(borders)

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(22)
    add_picture_with_alt(
        p.add_run(),
        LOGO,
        Inches(1.72),
        "Zap It logo: a yellow lightning bolt striking a target on a teal game tile.",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("ZAP IT")
    set_font(r, size=13, color=COLORS["accent"], bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Style Guide")
    set_font(r, size=34, color=COLORS["ink"], bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Visual language, interaction patterns, and implementation standards")
    set_font(r, size=13, color=COLORS["muted"])

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120], 0)
    for cell, fill in zip(table.rows[0].cells, [COLORS["yellow"], COLORS["accent"], COLORS["ink"]]):
        set_cell_shading(cell, fill)
        set_cell_margins(cell, top=75, bottom=75, start=0, end=0)

    add_text(
        doc,
        "Built from the production interface: logo artwork, CSS design tokens, game states, "
        "responsive behavior, themes, motion, and accessibility patterns.",
        size=10.2,
        color=COLORS["muted"],
        after=18,
        before=28,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_callout(
        doc,
        "Core idea",
        "Fast recognition, immediate feedback, and playful energy without sacrificing clarity.",
        fill="F7FAFB",
        accent=COLORS["yellow"],
    )
    add_text(
        doc,
        "Reference version 1.0  |  June 2026",
        size=8.8,
        color=COLORS["muted"],
        after=0,
        before=36,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_page_break()


def add_foundations(doc):
    add_kicker(doc, "01 / Brand foundations")
    add_heading(doc, "The Zap It design character", 1)
    add_text(
        doc,
        "Zap It is a focused reaction game. Its interface should feel quick, bright, tactile, and "
        "trustworthy. Every visual choice should help players identify the target, act confidently, "
        "and understand the result without slowing down.",
        size=10.5,
        after=10,
    )
    add_standard_table(
        doc,
        ["Principle", "What it means", "Design implication"],
        [
            ("Immediate", "The next action is always obvious.", "One dominant action; short labels; strong state changes."),
            ("Electric", "Energy comes from color, motion, and the bolt motif.", "Use yellow and zap effects as accents, not constant decoration."),
            ("Clear", "Speed never reduces legibility.", "High contrast, simple geometry, compact but breathable spacing."),
            ("Tactile", "Controls respond like physical game pieces.", "Pressed movement, impact feedback, depth through restrained shadows."),
            ("Inclusive", "Touch, keyboard, and assistive technology are first-class.", "Visible focus, semantic controls, live status, readable theme variants."),
        ],
        [1700, 3200, 4460],
    )

    add_heading(doc, "Experience priorities", 2)
    add_bullet(doc, "Recognition before decoration. The target and matching shape must dominate visual attention.", "Recognition")
    add_bullet(doc, "Feedback within a fraction of a second. Correct and incorrect actions must look unmistakably different.", "Feedback")
    add_bullet(doc, "A calm frame around an energetic game. Panels and typography stay restrained so gameplay can feel lively.", "A calm frame")
    add_bullet(doc, "Mobile first. The core play surface is designed for a narrow, one-handed viewport and expands modestly on larger screens.", "Mobile first")

    add_heading(doc, "Voice and terminology", 2)
    add_standard_table(
        doc,
        ["Use", "Avoid", "Reason"],
        [
            ("Start, Restart, Play again", "Initialize, Reset state", "Player language should describe the action, not the implementation."),
            ("Zap the matching item", "Select the correct option", "The brand verb gives instructions personality and specificity."),
            ("Wrong item. 10 seconds lost.", "Invalid input", "State the consequence plainly and immediately."),
            ("Zaps/sec", "Clicks per second", "Use the product's own vocabulary consistently."),
        ],
        [2600, 2600, 4160],
    )
    add_callout(doc, "Writing rule", "Keep interface copy active, concrete, and brief. Prefer one sentence over explanatory paragraphs.")
    doc.add_page_break()


def add_logo_and_shape(doc):
    add_kicker(doc, "02 / Identity")
    add_heading(doc, "Logo, icon, and shape language", 1)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [3000, 6360], 180)
    left, right = table.rows[0].cells
    set_cell_shading(left, "F7FAFB")
    set_cell_shading(right, "FFFFFF")
    for cell in (left, right):
        set_cell_border(cell, COLORS["line"], 7)
        set_cell_margins(cell, top=180, bottom=180, start=180, end=180)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_with_alt(
        p.add_run(),
        LOGO,
        Inches(1.55),
        "Zap It primary logo.",
    )
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Primary mark")
    set_font(r, size=12, color=COLORS["ink"], bold=True)
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(
        "A lightning bolt strikes a teal target tile. Supporting shapes communicate matching, "
        "speed, and playful competition."
    )
    set_font(r, size=9.5, color=COLORS["ink"])
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Use the full mark for app branding and the simplified icon for browser or compact contexts.")
    set_font(r, size=9.2, color=COLORS["muted"])

    add_heading(doc, "Usage rules", 2)
    add_standard_table(
        doc,
        ["Rule", "Specification"],
        [
            ("Clear space", "Keep at least 25% of the logo tile width clear on every side."),
            ("Minimum size", "Full logo: 40 px digital / 0.42 in print. Favicon: use the simplified icon artwork."),
            ("Background", "Prefer white, very light neutral, or the app's dark panel color. Preserve edge contrast."),
            ("Alteration", "Do not recolor individual logo elements, flatten the gradients, stretch, rotate, or add effects."),
        ],
        [2100, 7260],
    )

    add_heading(doc, "Game shapes", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    add_picture_with_alt(
        p.add_run(),
        SHAPES_IMAGE,
        Inches(6.25),
        "The five canonical Zap It shapes: star, circle, rounded square, triangle, and diamond.",
    )
    add_text(
        doc,
        "Star, circle, rounded square, triangle, and diamond are the canonical set. Shapes use a "
        "bold fill, a subtle dark outline, and enough interior scale to remain recognizable at speed.",
        size=9.7,
        color=COLORS["muted"],
        after=8,
    )
    add_callout(
        doc,
        "Shape rule",
        "Do not communicate a gameplay distinction with color alone. Shape and accessible labels remain part of the identity.",
    )
    doc.add_page_break()


def add_color(doc):
    add_kicker(doc, "03 / Color")
    add_heading(doc, "Core interface palette", 1)
    add_text(
        doc,
        "The core palette uses deep blue-black for clarity, teal for action, warm yellow for electric "
        "brand energy, and restrained semantic colors for outcomes.",
        after=8,
    )
    add_color_table(
        doc,
        [
            ("Ink", COLORS["ink"], "Primary text, dark outlines, high-contrast anchors"),
            ("Muted", COLORS["muted"], "Secondary copy, labels, supporting metrics"),
            ("Accent", COLORS["accent"], "Primary actions, timer ring, active emphasis"),
            ("Accent strong", COLORS["accent_strong"], "Icon color, links, small emphasis on light surfaces"),
            ("Panel", COLORS["panel"], "Primary surface and modal background"),
            ("Field", COLORS["field"], "Inset stat blocks, options, leaderboard rows"),
            ("Line", COLORS["line"], "Borders and quiet structural dividers"),
            ("Electric yellow", COLORS["yellow"], "Logo bolt and selective brand highlights"),
            ("Success", COLORS["good"], "Correct match feedback"),
            ("Danger", COLORS["danger"], "Miss feedback and time penalty"),
        ],
    )

    add_heading(doc, "Color hierarchy", 2)
    add_bullet(doc, "Teal is functional. Reserve it for actions, focus, timing, and active structure.", "Teal")
    add_bullet(doc, "Yellow is expressive. Use it sparingly for identity and moments of energy.", "Yellow")
    add_bullet(doc, "Green and magenta-red are semantic. Never reuse them casually where correctness could be inferred.", "Green and magenta-red")
    add_bullet(doc, "Neutral surfaces carry most of the interface. This keeps the colored shapes visually dominant.", "Neutral surfaces")

    add_heading(doc, "Dark mode", 2)
    add_color_table(
        doc,
        [
            ("Dark panel", COLORS["dark_panel"], "Main surface"),
            ("Dark field", COLORS["dark_field"], "Inset controls and stat regions"),
            ("Dark line", COLORS["dark_line"], "Borders"),
            ("Dark accent", "40B9C8", "Primary actions and timer ring"),
            ("Dark ink", "F1F7FA", "Primary text"),
            ("Dark muted", "A9BBC4", "Secondary text"),
        ],
    )
    add_callout(
        doc,
        "Contrast",
        "Treat light and dark mode as separate tested palettes. Do not create dark mode by simply inverting colors.",
        fill="E8F7F9",
        accent=COLORS["accent"],
    )
    doc.add_page_break()


def add_themes(doc):
    add_kicker(doc, "04 / Themes")
    add_heading(doc, "Theme palettes", 1)
    add_text(
        doc,
        "Themes change the atmosphere and shape palette while preserving the same hierarchy, interaction "
        "states, component geometry, and accessibility expectations.",
        after=9,
    )
    add_standard_table(
        doc,
        ["Theme", "Primary colors", "Character", "Usage note"],
        [
            ("Normal", "#2877D8  #22A06B  #D83B4A  #F0C83B  #EF8737", "Balanced and familiar", "Default multicolor game experience"),
            ("One color", "User-selected base plus light/dark mixes", "Focused and customizable", "Shapes use gradients; tiles receive subtle tonal variation"),
            ("Christmas", "#C8D0D6  #C52B2F  #167A4A", "Festive and crisp", "Keep silver readable against both surface modes"),
            ("Fall", "#C65F24  #D99A22  #A83F3F  #6F7F3A  #7B4B5B", "Warm and earthy", "Use cream-toned ambient backgrounds"),
            ("Summer", "#F8C630  #FF6F61  #24B8C8  #8CC63F  #4F9FEA", "Bright and buoyant", "Maintain strong borders around light fills"),
        ],
        [1200, 3300, 1900, 2960],
        font_size=8.2,
    )
    add_heading(doc, "Theme construction rules", 2)
    add_bullet(doc, "Keep the base neutral system stable: text, panels, fields, lines, focus, and semantic feedback must remain predictable.")
    add_bullet(doc, "Change ambient background gradients and accent values, but preserve panel readability and action prominence.")
    add_bullet(doc, "In one-color mode, calculate a light mix and dark mix from the selected color, then choose button text by luminance.")
    add_bullet(doc, "Test every theme in light and dark mode, during gameplay and inside modal screens.")

    add_heading(doc, "One-color recipe", 2)
    add_standard_table(
        doc,
        ["Output", "Mix", "Current implementation"],
        [
            ("Light companion", "Base toward white", "42% white for theme highlight; 72-88% white for tiles"),
            ("Dark companion", "Base toward black", "20% black for accent; 24% black for shape gradient"),
            ("Button text", "Choose by luminance", "Dark ink above a 0.62 luminance threshold; white otherwise"),
        ],
        [2100, 2600, 4660],
    )
    add_callout(
        doc,
        "Guardrail",
        "Seasonal personality belongs in color and atmosphere. Do not change gameplay rules, controls, or information architecture by theme.",
    )
    doc.add_page_break()


def add_type_layout(doc):
    add_kicker(doc, "05 / Typography and layout")
    add_heading(doc, "Type system", 1)
    add_text(
        doc,
        "Zap It uses Arial with Helvetica and system sans-serif fallbacks. The system favors familiar "
        "letterforms, compact hierarchy, and heavy weights for rapid scanning.",
        after=8,
    )
    add_standard_table(
        doc,
        ["Role", "App size", "Weight", "Treatment"],
        [
            ("Primary app heading", "1.32rem", "Bold", "Tight line height; sentence case"),
            ("Modal heading", "1.45-2.1rem", "Bold", "Short phrases; strong contrast"),
            ("Target name", "1.08rem", "Extra bold", "Centered; immediate gameplay priority"),
            ("Eyebrow / label", "0.72rem", "Bold", "Uppercase; muted; compact"),
            ("Body / instructions", "0.86-0.92rem", "Regular", "1.18-1.3 line height"),
            ("Metrics", "0.95-1.25rem", "Extra bold", "Tabular-feeling alignment; short labels"),
            ("Legal / disclosure", "0.68-0.76rem", "Regular", "Use only for genuinely secondary required text"),
        ],
        [2500, 1600, 1500, 3760],
    )

    add_heading(doc, "Spacing and geometry", 2)
    add_standard_table(
        doc,
        ["Token", "Value", "Use"],
        [
            ("Core radius", "0.5rem", "Panels, buttons, cells, dialogs, options"),
            ("Compact radius", "0.45rem", "Inset fields, rows, secondary controls"),
            ("Large shell radius", "0.75rem", "Desktop game shell"),
            ("Core gap", "0.4-0.75rem", "Primary layout rhythm"),
            ("Touch control", "2.55-2.85rem min", "Icon buttons, inputs, settings rows"),
            ("Play grid", "3 x 3, square", "Maximum 25rem mobile / 27rem larger viewports"),
            ("Content shell", "30rem max", "Narrow, phone-first centered game experience"),
        ],
        [2400, 1900, 5060],
    )

    add_heading(doc, "Surface and depth", 2)
    add_bullet(doc, "Panels use a soft translucent white surface in light mode and an opaque dark surface in dark mode.", "Panels")
    add_bullet(doc, "Borders define structure before shadows do. Keep the line color quiet but visible.", "Borders")
    add_bullet(doc, "Use soft elevation for panels and stronger elevation for primary actions and overlays.", "Elevation")
    add_bullet(doc, "The game shapes carry their own drop shadow so they separate from the tile without making the tile visually noisy.", "Shapes")
    add_callout(
        doc,
        "Layout rule",
        "Keep the main play loop visible in one narrow column. Avoid sidebars or persistent secondary navigation near the grid.",
    )
    doc.add_page_break()


def add_components(doc):
    add_kicker(doc, "06 / Components")
    add_heading(doc, "Core component patterns", 1)
    add_standard_table(
        doc,
        ["Component", "Visual recipe", "Behavior"],
        [
            ("Top bar", "Soft panel, logo lockup, compact icon actions", "Keeps brand and utility controls available without competing with play"),
            ("Target panel", "Centered target name and shape; circular timer at right", "Waiting state removes shape and timer to simplify the start moment"),
            ("Stat row", "Three equal inset fields", "Score, streak, and rate update continuously"),
            ("Grid cell", "2 px border, 0.5rem radius, pale gradient, centered shape", "Press scales to 96%; focus gets a 3 px outline"),
            ("Primary button", "Solid accent fill, white or computed readable text", "One dominant action per view; 1 px pressed movement"),
            ("Ghost button", "Panel fill, line border, strong accent text", "Secondary or refresh actions"),
            ("Overlay / dialog", "58% dark scrim, compact panel, strong shadow", "Focus moves into the dialog and returns to the trigger on close"),
            ("Leaderboard row", "Field background, rank accent, compact metrics", "Current run receives an inset accent ring"),
        ],
        [1700, 4080, 3580],
        font_size=8.15,
    )

    add_heading(doc, "Button hierarchy", 2)
    add_standard_table(
        doc,
        ["Level", "Use", "Examples"],
        [
            ("Primary", "Advance or restart the main flow", "Start, Save, Play again"),
            ("Secondary", "Optional maintenance or refresh action", "Refresh rankings"),
            ("Icon", "Compact utility action with a clear accessible name", "How to play, Leaderboard, Settings, Close"),
        ],
        [1700, 4500, 3160],
    )

    add_heading(doc, "State styling", 2)
    add_standard_table(
        doc,
        ["State", "Required signal", "Current treatment"],
        [
            ("Default", "Clear affordance", "Border, surface contrast, pointer cursor"),
            ("Pressed", "Physical response", "Translate 1 px or scale to 96%"),
            ("Focus visible", "Keyboard location", "3 px high-contrast outline with 3 px offset"),
            ("Disabled", "Unavailable but legible", "62% opacity for controls; 55% for grid cells"),
            ("Correct", "Positive outcome", "Green border and soft outer ring plus zap impact"),
            ("Wrong", "Error and consequence", "Danger border, soft ring, horizontal shake, timer penalty pulse"),
        ],
        [1800, 3600, 3960],
    )
    add_callout(
        doc,
        "Component rule",
        "A state must remain understandable without animation. Motion reinforces the border, color, copy, and focus change.",
    )
    doc.add_page_break()


def add_motion_accessibility(doc):
    add_kicker(doc, "07 / Motion and accessibility")
    add_heading(doc, "Motion language", 1)
    add_text(
        doc,
        "Motion is short, responsive, and tied to cause and effect. The interface should never animate "
        "merely to stay busy.",
        after=8,
    )
    add_standard_table(
        doc,
        ["Motion", "Duration", "Purpose", "Character"],
        [
            ("Lightning strike", "140 ms", "Confirm a correct zap", "Sharp, stepped flash with glow and impact point"),
            ("Impact compression", "130 ms", "Make the chosen tile feel struck", "Quick scale and brightness pulse"),
            ("Wrong shake", "190 ms", "Signal a miss", "Small horizontal displacement"),
            ("Timer penalty", "420 ms", "Connect the miss to lost time", "Danger color and brief scale increase"),
            ("Board refresh", "520 ms", "Separate rounds and reveal the next set", "Fast 3D flip with a mid-animation content swap"),
        ],
        [2100, 1200, 3000, 3060],
    )
    add_bullet(doc, "Keep gameplay feedback below roughly half a second so it feels immediate.")
    add_bullet(doc, "Lock the board during a round transition to prevent accidental double input.")
    add_bullet(doc, "Avoid springy or decorative motion that obscures shape recognition.")
    add_bullet(doc, "Provide a reduced-motion path when the product adds broader motion preferences.")

    add_heading(doc, "Accessibility requirements", 2)
    add_standard_table(
        doc,
        ["Area", "Standard"],
        [
            ("Semantics", "Use native buttons, fieldsets, labels, ordered lists, and dialog roles."),
            ("Names", "Every icon-only control needs an accessible name and matching title where useful."),
            ("Status", "Timer, messages, score-save feedback, and leaderboard loading use appropriate live regions."),
            ("Keyboard", "Numpad mirrors the 3 x 3 board; Space restarts; Escape closes overlays."),
            ("Focus", "Move focus into opened dialogs and return it to the originating control on close."),
            ("Contrast", "Test each seasonal palette in both modes; retain borders around pale shapes."),
            ("Privacy", "Keep the score-sharing warning visible and avoid requesting identifying information."),
        ],
        [1900, 7460],
    )
    add_callout(
        doc,
        "Non-negotiable",
        "Do not remove visible focus, accessible labels, dialog focus management, or live feedback while restyling components.",
        fill="FFF4F6",
        accent=COLORS["danger"],
    )
    doc.add_page_break()


def add_responsive_checklist(doc):
    add_kicker(doc, "08 / Responsive and implementation")
    add_heading(doc, "Responsive behavior", 1)
    add_standard_table(
        doc,
        ["Viewport", "Behavior"],
        [
            ("Base / mobile", "Full-height narrow shell, 0.7rem outer padding, 25rem maximum grid, touch-first controls."),
            ("42rem and wider", "Center the app in the viewport, add body padding, increase shell padding, allow a 27rem grid."),
            ("24rem and narrower", "Hide leaderboard rate details and stack the score-save input and button."),
            ("Active play", "Prevent body scrolling and overscroll so the game surface remains stable."),
        ],
        [2200, 7160],
    )

    add_heading(doc, "Implementation checklist", 2)
    checklist = [
        ("Brand", "Use the approved logo artwork and canonical five-shape set."),
        ("Tokens", "Reference shared variables for ink, muted, surfaces, line, accent, success, danger, focus, and elevation."),
        ("Hierarchy", "Keep one primary action and one dominant gameplay target per view."),
        ("Themes", "Verify Normal, One color, Christmas, Fall, and Summer in light and dark mode."),
        ("States", "Test default, pressed, focus, disabled, correct, wrong, loading, empty, and saved states."),
        ("Motion", "Preserve cause-and-effect timing and board locking during transitions."),
        ("Responsive", "Check narrow phone, standard phone, and centered desktop layouts."),
        ("Accessibility", "Retain semantic HTML, accessible names, focus return, keyboard mapping, and live status."),
        ("Privacy", "Keep the public-score warning and player-name validation intact."),
        ("Copy", "Use short player-facing language and the verb “zap” consistently."),
    ]
    add_standard_table(doc, ["Check", "Requirement"], checklist, [1700, 7660], font_size=8.65)

    add_heading(doc, "Do / do not", 2)
    add_standard_table(
        doc,
        ["Do", "Do not"],
        [
            ("Use color to establish action and outcome.", "Color every surface or turn the interface into a palette showcase."),
            ("Keep geometry simple and repeated.", "Introduce unrelated radii, ornate icons, or decorative textures."),
            ("Make feedback immediate and redundant.", "Depend on animation or color alone."),
            ("Let the shapes be the playful focal point.", "Add visual noise around the grid."),
            ("Write for a player in motion.", "Expose technical or backend language in the interface."),
        ],
        [4680, 4680],
    )

    add_callout(
        doc,
        "Source of truth",
        "This guide reflects the current production files: assets/brand/logo.svg, styles.css, index.html, and app.js.",
        fill="F7FAFB",
        accent=COLORS["yellow"],
    )
    add_text(
        doc,
        "End of guide",
        size=8.5,
        color=COLORS["muted"],
        after=0,
        before=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    create_shape_image()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_foundations(doc)
    add_logo_and_shape(doc)
    add_color(doc)
    add_themes(doc)
    add_type_layout(doc)
    add_components(doc)
    add_motion_accessibility(doc)
    add_responsive_checklist(doc)
    core = doc.core_properties
    core.title = "Zap It Style Guide"
    core.subject = "Visual language, interaction patterns, and implementation standards"
    core.author = "Zap It"
    core.keywords = "Zap It, style guide, design system, brand, UI"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
